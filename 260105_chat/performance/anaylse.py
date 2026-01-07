import re
import pandas as pd
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def parse_performance_log(file_path):
    """
    解析性能日志文件，提取每个 token 的性能数据
    """
    data = []
    
    with open(file_path, 'r') as f:
        lines = f.readlines()
    
    # 定义正则表达式匹配模式
    # 格式: [prof] token <pos> <name> <dims>: <cycles> cycles (<ms> ms)
    # 支持负数 token_pos（如 -1 表示 encode）以及字母字符串
    pattern = r'\[prof\] token (-?\d+|[\w]+)\s+([\w_]+)\s*(?:\(([\d,]+)\))?\s*:\s*(\d+)\s+cycles\s+\((\d+(?:\.\d+)?)\s+ms\)'
    
    # 辅助变量
    current_token = None
    current_operation_data = {}
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
        
        # 尝试匹配性能数据行
        match = re.search(pattern, line)
        if match:
            token_pos = match.group(1)
            operation_name = match.group(2)
            dims = match.group(3) if match.group(3) else ''
            cycles = int(match.group(4))
            ms = float(match.group(5))
            
            # 处理特殊的 token 值 (如 18446744073709551615)
            if token_pos == '18446744073709551615':
                token_pos = -1
            
            # 添加到数据
            data.append({
                'token_pos': token_pos,
                'operation': operation_name,
                'dims': dims,
                'cycles': cycles,
                'ms': ms
            })
        
        # 解析总统计信息
        elif '[bare]' in line:
            if 'tok/s:' in line:
                tok_per_sec = float(line.split(':')[1].strip())
                data.append({
                    'token_pos': 'SUMMARY',
                    'operation': 'tokens_per_second',
                    'dims': '',
                    'cycles': 0,
                    'ms': tok_per_sec
                })
            elif 'total time/ms:' in line:
                total_ms = int(line.split(':')[1].strip())
                data.append({
                    'token_pos': 'SUMMARY',
                    'operation': 'total_time_ms',
                    'dims': '',
                    'cycles': 0,
                    'ms': total_ms
                })
            elif 'TTFT:' in line:
                ttft_ms = int(line.split(':')[1].replace('ms', '').strip())
                data.append({
                    'token_pos': 'SUMMARY',
                    'operation': 'TTFT_ms',
                    'dims': '',
                    'cycles': 0,
                    'ms': ttft_ms
                })
    
    return data

def create_summary_statistics(data):
    """
    创建汇总统计信息
    """
    # 过滤出非SUMMARY数据
    perf_data = [d for d in data if d['token_pos'] != 'SUMMARY']
    
    if not perf_data:
        return pd.DataFrame()
    
    df = pd.DataFrame(perf_data)
    
    # 转换token_pos为数字
    def to_numeric(x):
        try:
            return int(x)
        except:
            return -2  # 给非数字token_pos一个特殊值
    
    df['token_num'] = df['token_pos'].apply(to_numeric)
    
    # 按操作类型汇总
    operation_summary = df.groupby('operation').agg({
        'cycles': ['sum', 'mean', 'min', 'max', 'count'],
        'ms': ['sum', 'mean', 'min', 'max']
    }).round(2)
    
    # 扁平化列名
    operation_summary.columns = [f'{col[0]}_{col[1]}' for col in operation_summary.columns]
    operation_summary = operation_summary.reset_index()
    
    # 按token位置汇总
    token_summary = df[df['token_num'] >= 0].groupby('token_num').agg({
        'cycles': 'sum',
        'ms': 'sum'
    }).round(2).reset_index()
    
    # 添加每token的平均时间
    if len(token_summary) > 0:
        avg_per_token = token_summary['ms'].mean()
        token_summary['avg_ms_per_token'] = avg_per_token
    
    # 获取SUMMARY数据
    summary_data = {d['operation']: d['ms'] for d in data if d['token_pos'] == 'SUMMARY'}
    
    return operation_summary, token_summary, summary_data

def save_to_excel(data, operation_summary, token_summary, summary_data, output_path):
    """
    将数据保存到Excel文件
    """
    # 创建DataFrame
    df_all = pd.DataFrame(data)
    
    # 创建Excel写入器
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        # 1. 原始数据工作表
        df_all.to_excel(writer, sheet_name='Raw Data', index=False)
        
        # 2. 操作汇总工作表
        if not operation_summary.empty:
            operation_summary.to_excel(writer, sheet_name='Operation Summary', index=False)
        
        # 3. Token汇总工作表
        if not token_summary.empty:
            token_summary.to_excel(writer, sheet_name='Token Summary', index=False)
        
        # 4. 整体统计工作表
        summary_df = pd.DataFrame(list(summary_data.items()), columns=['Metric', 'Value'])
        summary_df.to_excel(writer, sheet_name='Overall Statistics', index=False)
    
    # 应用格式
    apply_excel_formatting(output_path)

def apply_excel_formatting(file_path):
    """
    应用Excel格式
    """
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    
    wb = load_workbook(file_path)
    
    # 定义样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    center_alignment = Alignment(horizontal="center")
    
    # 应用到每个工作表
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # 设置列宽
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # 设置标题行样式
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_alignment
        
        # 冻结首行
        ws.freeze_panes = 'A2'
    
    wb.save(file_path)

def main():
    # 输入文件路径
    input_file = 'performance_layer_debug4.txt'
    
    # 检查文件是否存在
    if not os.path.exists(input_file):
        print(f"错误: 文件 {input_file} 不存在!")
        return
    
    # 输出文件路径
    output_file = 'performance_analysis_matrix_debug4.xlsx'
    
    print("开始解析性能日志文件...")
    
    # 解析数据
    data = parse_performance_log(input_file)
    
    if not data:
        print("未找到有效数据!")
        return
    
    print(f"成功解析 {len(data)} 行数据")
    
    # 创建汇总统计
    operation_summary, token_summary, summary_data = create_summary_statistics(data)
    
    # 保存到Excel
    save_to_excel(data, operation_summary, token_summary, summary_data, output_file)
    
    print(f"分析完成! 结果已保存到: {output_file}")
    
    # 打印关键统计信息
    print("\n关键统计信息:")
    print("-" * 50)
    
    if 'tokens_per_second' in summary_data:
        print(f"Tokens per second: {summary_data['tokens_per_second']:.4f}")
    
    if 'total_time_ms' in summary_data:
        print(f"Total time: {summary_data['total_time_ms'] / 1000:.2f} seconds")
    
    if 'TTFT_ms' in summary_data:
        print(f"TTFT: {summary_data['TTFT_ms'] / 1000:.2f} seconds")
    
    if not token_summary.empty and 'avg_ms_per_token' in token_summary.columns:
        print(f"Average time per token: {token_summary['avg_ms_per_token'].iloc[0]:.2f} ms")
    
    if not operation_summary.empty:
        print(f"\nTop 5 most expensive operations (by total cycles):")
        top_ops = operation_summary.nlargest(5, 'cycles_sum')
        for _, row in top_ops.iterrows():
            print(f"  {row['operation']}: {int(row['cycles_sum'])} cycles, {row['ms_sum']:.2f} ms")

if __name__ == "__main__":
    main()